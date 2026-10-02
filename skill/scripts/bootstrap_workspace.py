# -*- coding: utf-8 -*-
"""初始化新工作区并安装 skills —— 在解压后的 bundle 内运行

用法:
    python bootstrap_workspace.py <目标工作区路径> [--force]

不做的事: 不安装系统依赖（先跑 setup_env.py）；不迁移旧数据（旧机的
weekly/、posts/、topics.xlsx 手动拷到新工作区同名位置即可）。
"""
import os
import shutil
import sys

HERE = os.path.dirname(os.path.abspath(__file__))          # .../bundle/skill/scripts
BUNDLE = os.path.dirname(os.path.dirname(HERE))            # .../bundle
SEED = os.path.join(BUNDLE, "workspace-seed")
HUMANIZE = os.path.join(BUNDLE, "humanize-writing")


def copy_tree(src, dst, force=False):
    for root, _dirs, files in os.walk(src):
        rel = os.path.relpath(root, src)
        troot = os.path.join(dst, rel) if rel != "." else dst
        os.makedirs(troot, exist_ok=True)
        for f in files:
            tf = os.path.join(troot, f)
            if os.path.exists(tf) and not force:
                print(f"  跳过已存在: {os.path.relpath(tf, dst)}")
                continue
            shutil.copy2(os.path.join(root, f), tf)


def build_topics(path):
    """自包含构建 topics.xlsx（不依赖 xlsx skill 的 base.py）"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    FONT = "Microsoft YaHei"
    hdr_fill = PatternFill("solid", fgColor="1B2A4A")
    hdr_font = Font(name=FONT, size=11, color="FFFFFF")
    body = Font(name=FONT, size=11, color="37352F")
    alt = PatternFill("solid", fgColor="F7F7F5")
    hdr_border = Border(bottom=Side(style="thin", color="E9E9E8"))
    wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def sheet(wb_, name, title, headers, notes, widths):
        ws = wb_.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = 3
        ws.row_dimensions[1].height = 15
        ws.row_dimensions[2].height = 32
        ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=1 + len(headers))
        t = ws.cell(row=2, column=2, value=title)
        t.font = Font(name=FONT, size=16, color="1B2A4A")
        t.alignment = Alignment(horizontal="left", vertical="center")
        for i, h in enumerate(headers):
            c = ws.cell(row=4, column=2 + i, value=h)
            c.fill, c.font, c.border = hdr_fill, hdr_font, hdr_border
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[4].height = 28
        for i, n in enumerate(notes):
            c = ws.cell(row=6 + len(notes) + i + 1, column=2, value=n)
            c.font = Font(name=FONT, size=9, color="8C8A84")
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(ord("B") + i)].width = w
        return ws

    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, "候选清单", "选题候选清单",
          ["周次", "日期", "标题", "来源", "链接", "一句话理由", "热度源", "初评", "状态"],
          ["说明：采集运行自动/手动追加；周次与 HF 周榜 URL（week/<年份>-Wxx）对齐。",
           "初评=热度×题材适配度；状态：候选 / 已选 / 成稿 / 已发布 / 弃。"],
          [10, 12, 34, 14, 30, 34, 8, 8, 10])
    sheet(wb, "发布记录", "发布记录与数据复盘",
          ["发布日期", "笔记标题", "关联选题", "平台", "笔记链接", "浏览", "点赞", "收藏", "评论", "复盘备注"],
          ["说明：发布 48-72 小时后回填数据；复盘备注写「哪句钩子有效/哪个标签带量」。"],
          [12, 30, 26, 10, 26, 8, 8, 8, 8, 24])
    wb.properties.creator = "Z.ai"
    wb.save(path)


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    force = "--force" in sys.argv
    if not args:
        sys.exit(__doc__)
    target = os.path.abspath(args[0])
    if os.path.exists(os.path.join(target, "sources.md")) and not force:
        sys.exit(f"目标已有 sources.md（{target}）。确认覆盖请加 --force")

    print(f"初始化工作区: {target}\n")
    os.makedirs(target, exist_ok=True)

    # 1. 工作区种子
    copy_tree(SEED, target, force)
    print("  种子文件已复制（sources.md、templates/）")

    # 2. 目录骨架
    for d in ("weekly", "posts"):
        os.makedirs(os.path.join(target, d), exist_ok=True)
    print("  目录骨架已创建（weekly/、posts/）")

    # 3. 台账
    tp = os.path.join(target, "topics.xlsx")
    if not os.path.exists(tp) or force:
        build_topics(tp); print("  topics.xlsx 已生成")
    else:
        print("  跳过已存在: topics.xlsx")

    # 4. 安装 skills
    skill_dst = os.path.join(target, ".agents", "skills", "lunjian")
    copy_tree(os.path.join(BUNDLE, "skill"), skill_dst, force)
    print(f"  lunjian skill → {skill_dst}")
    if os.path.isdir(HUMANIZE):
        hu_dst = os.path.join(os.path.expanduser("~"), ".agents", "skills", "humanize-writing")
        copy_tree(HUMANIZE, hu_dst, force)
        print(f"  humanize-writing skill → {hu_dst}（用户级，跨项目可用）")
    else:
        print("  ⚠️ bundle 内未找到 humanize-writing/，去 AI 味工序将不可用（可从 github.com/jpeggdev/humanize-writing 获取）")

    print("""
完成。下一步:
  1. python <bundle>/skill/scripts/setup_env.py   # 检查/安装依赖
  2. 用 ZCode 打开该工作区，说「跑今天」或「解读这篇 + 链接」
  3. 旧数据迁移: 把旧机 weekly/、posts/、topics.xlsx 拷到新工作区同名位置
""")


if __name__ == "__main__":
    main()
